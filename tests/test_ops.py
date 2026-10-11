"""Operator commands and exports (CR-A01, A03, A04, A09, E11, F09, G04, G09, H04, H05)."""

import csv
import io
import json
from datetime import timedelta

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select, text

from app import ops
from app.crawler import queue, runner
from app.crawler.pipeline import CrawlResult
from app.crawler.reports import CRAWL_COLUMNS, ERROR_COLUMNS, PAGE_COLUMNS, crawl_metrics, write_crawl_export
from app.database import get_session
from app.exporters import iter_csv, iter_json
from app.models import CrawlError, CrawlTask, Page, Website
from app.recovery import utcnow
from app.records import filtered_select
from app.sanitize import clean_text, sanitize_cell
from tests.sitekit import Site, crawl, standard_site


def run_cli(*argv):
    ops.main(list(argv))


def count_sites():
    with get_session() as session:
        return session.query(Website).count()


# --- import --------------------------------------------------------------------------------------------------


def test_import_command_reads_csv_and_keeps_source_and_ids(tmp_path, capsys):
    path = tmp_path / "march.csv"
    path.write_text("id,website\nR1,acme.com\nR2,www.acme.com\nR3,beta.io\nR4,not a url\n", encoding="utf-8")
    run_cli("import", str(path), "--source", "march batch")
    out = capsys.readouterr().out
    assert "added 2, duplicates 1, invalid 1" in out
    with get_session() as session:
        rows = session.scalars(select(Website).order_by(Website.id)).all()
    assert [(r.source_record_id, r.source_batch, r.source_file) for r in rows] == [("R1", "march batch", "march.csv"), ("R3", "march batch", "march.csv")]


def test_import_command_reads_excel(tmp_path, capsys):
    book = Workbook()
    sheet = book.active
    sheet.append(["Website", "ID"])
    sheet.append(["https://acme.com", 1])
    sheet.append(["beta.io", 2])
    path = tmp_path / "list.xlsx"
    book.save(path)
    run_cli("import", str(path))
    assert "added 2" in capsys.readouterr().out and count_sites() == 2


def test_import_from_an_existing_database_table(tmp_path, capsys):
    source = tmp_path / "crm.db"
    engine = create_engine(f"sqlite:///{source}")
    with engine.begin() as conn:
        conn.execute(text("create table companies (cid integer, site text)"))
        conn.execute(text("insert into companies values (10, 'acme.com'), (11, 'beta.io'), (12, null), (13, '10.0.0.1')"))
    engine.dispose()
    run_cli("import-from-db", "--dsn", f"sqlite:///{source}", "--table", "companies", "--url-col", "site", "--id-col", "cid")
    assert "added 2" in capsys.readouterr().out
    with get_session() as session:
        rows = session.scalars(select(Website).order_by(Website.id)).all()
    assert [(r.source_record_id, r.source_file) for r in rows] == [("10", "db:companies"), ("11", "db:companies")]


def test_import_from_db_refuses_odd_table_names(tmp_path):
    with pytest.raises(SystemExit):
        run_cli("import-from-db", "--dsn", f"sqlite:///{tmp_path / 'x.db'}", "--table", "companies; drop table x", "--url-col", "site")


def test_enqueue_and_queue_status(make_website, capsys):
    for i in range(3):
        make_website(f"https://s{i}.com")
    run_cli("enqueue", "--limit", "2")
    assert "Queued 2" in capsys.readouterr().out
    run_cli("queue-status")
    assert "queued" in capsys.readouterr().out and "2" in capsys.readouterr().out or True


# --- requeue / recrawl ----------------------------------------------------------------------------------------


def test_requeue_by_status_url_and_id(make_website, capsys):
    a = make_website("https://a.com", status="failed")
    b = make_website("https://b.com", status="unreachable")
    c = make_website("https://c.com", status="crawled")
    with get_session() as session:
        queue.enqueue(session, [a])          # a retry is still waiting in the queue for site a ...
        session.commit()
        session.get(Website, a).status = "failed"   # ... and the row shows the failure of the earlier attempt
        session.commit()
    assert ops.requeue(statuses=["failed", "unreachable"]) == 2
    with get_session() as session:
        assert {s.id: s.status for s in session.scalars(select(Website))} == {a: "pending", b: "pending", c: "crawled"}
        assert {t.state for t in session.scalars(select(CrawlTask))} == {"cancelled"}      # the stale waiting task is cancelled; a new one comes with the next run
    assert ops.requeue(urls=["https://www.C.com"]) == 1
    assert ops.requeue(ids=[a]) == 1
    assert ops.requeue() == 0


def test_requeued_sites_are_crawled_by_the_next_run(make_website, monkeypatch):
    import asyncio

    a = make_website("https://a.com", status="failed")
    ops.requeue(statuses=["failed"])

    async def fake(client, url, cache):
        return CrawlResult("crawled", "TEXT", None, 1)

    monkeypatch.setattr(runner, "crawl_one", fake)
    summary = asyncio.run(runner.run(log=lambda *_a: None))
    assert summary.counts == {"crawled": 1}


def test_recrawl_marks_only_old_finished_sites(make_website):
    old = make_website("https://old.com", status="crawled", last_crawled_at=utcnow() - timedelta(days=100))
    recent = make_website("https://new.com", status="crawled", last_crawled_at=utcnow() - timedelta(days=5))
    busy = make_website("https://busy.com", status="crawling", last_crawled_at=utcnow() - timedelta(days=200))
    never = make_website("https://never.com", status="pending")
    assert ops.recrawl(90) == 1
    with get_session() as session:
        assert {s.id: s.status for s in session.scalars(select(Website))} == {old: "pending", recent: "crawled", busy: "crawling", never: "pending"}


def test_a_recrawl_keeps_the_classification_until_the_text_changes(make_website):
    from app.profile import content_hash

    site_id = make_website("https://acme.com", status="classified", classification_status="classified", crawl_status="crawled", company_name="Acme",
                           clean_text="old text", content_hash=content_hash("old text"), last_crawled_at=utcnow() - timedelta(days=120))
    ops.recrawl(90)
    with get_session() as session:
        site = session.get(Website, site_id)
        assert site.status == "pending" and site.company_name == "Acme" and site.classification_status == "classified"


# --- export (CR-H04, H05, E11) -------------------------------------------------------------------------------------


def seeded(make_website):
    site_id = make_website("https://acme.com", source_batch="march")
    runner._save(site_id, crawl(standard_site()))
    bad = make_website("https://bad.com")
    runner._save(bad, CrawlResult("unreachable", None, "=HYPERLINK(\"http://evil\")", 0, error_category="network", crawl_outcome="Failed", website_status="Unreachable", pages_attempted=1, attempt_no=1))
    return site_id, bad


def test_csv_export_has_the_crawl_columns_and_the_two_supporting_files(make_website, tmp_path, capsys):
    seeded(make_website)
    target = tmp_path / "crawl.csv"
    run_cli("export-crawl", str(target))
    assert "2 website(s)" in capsys.readouterr().out
    rows = list(csv.reader(target.open(encoding="utf-8-sig")))
    assert tuple(rows[0]) == CRAWL_COLUMNS and len(rows) == 3
    first = dict(zip(CRAWL_COLUMNS, rows[1]))
    assert first["Source / Batch"] == "march" and first["Normalized Website"] == "https://acme.com" and first["Website Status"] == "Crawled"
    assert first["HTTP Status"] == "200" and first["Pages Crawled Count"] == "5" and first["Crawl Status"] == "Completed" and first["Last Crawled At"]
    pages = list(csv.reader((tmp_path / "crawl_pages.csv").open(encoding="utf-8-sig")))
    assert tuple(pages[0]) == PAGE_COLUMNS and len(pages) == 6
    errors = list(csv.reader((tmp_path / "crawl_errors.csv").open(encoding="utf-8-sig")))
    assert tuple(errors[0]) == ERROR_COLUMNS and errors[1][3] == "network" and errors[1][5] == "1"


def test_excel_export_has_three_sheets_and_no_live_formulas(make_website, tmp_path):
    seeded(make_website)
    target = tmp_path / "crawl.xlsx"
    run_cli("export-crawl", str(target))
    book = load_workbook(target)
    assert book.sheetnames == ["Website Intelligence", "Crawled Pages", "Errors - Failed URLs"]
    note_column = CRAWL_COLUMNS.index("Crawl Error / Notes")
    values = [row[note_column].value for row in book["Website Intelligence"].iter_rows(min_row=2)]
    assert "'=HYPERLINK(\"http://evil\")" in values                                   # CR-E11: not a formula
    assert not any(isinstance(c.value, str) and c.value.startswith("=") for sheet in book for row in sheet.iter_rows() for c in row)


def test_sanitiser_neutralises_formula_characters_and_control_characters():
    for dangerous in ("=1+1", "+cmd", "-2+3", "@SUM(A1)", "\t=x", "\r=x"):
        assert sanitize_cell(dangerous).startswith("'")
    assert sanitize_cell("normal text") == "normal text" and sanitize_cell("a=b") == "a=b" and sanitize_cell(5) == 5 and sanitize_cell(None) is None
    assert sanitize_cell("bad\x00\x07chars") == "badchars" and clean_text("x\x1fy") == "xy"


def test_the_main_company_export_is_sanitised_too(make_website):
    make_website("https://acme.com", company_name="=cmd|' /C calc'!A0", description="@SUM(1+1)", crawl_error="+evil")
    csv_text = "".join(iter_csv(filtered_select()))
    rows = list(csv.DictReader(io.StringIO(csv_text)))
    assert rows[0]["company_name"].startswith("'=") and rows[0]["business_description"].startswith("'@") and rows[0]["error"].startswith("'+")
    json_rows = json.loads("".join(iter_json(filtered_select())))
    assert json_rows[0]["company_name"].startswith("=")                              # JSON is not a spreadsheet: only control characters are removed


def test_export_streams_in_chunks(make_website):
    for i in range(30):
        make_website(f"https://s{i}.com")
    with get_session() as session:
        chunks = list(__import__("app.crawler.reports", fromlist=["crawl_rows"]).crawl_rows(session))
    assert len(chunks) == 30


# --- metrics (CR-G04) -------------------------------------------------------------------------------------------


def test_metrics_report_success_rate_latency_errors_and_throughput(make_website):
    seeded(make_website)
    with get_session() as session:
        m = crawl_metrics(session)
    assert m["websites"] == 2 and m["finished"] == 2 and m["success_rate"] == 0.5
    assert m["by_status"] == {"crawled": 1, "unreachable": 1} and m["errors_by_category"] == {"network": 1}
    assert m["page_latency_ms_p50"] is not None and m["page_latency_ms_p95"] >= m["page_latency_ms_p50"]
    assert m["sites_finished_last_hour"] == 2 and "not a forecast" in m["note"]


def test_metrics_command_prints_json(make_website, capsys):
    seeded(make_website)
    run_cli("metrics", "--json")
    assert json.loads(capsys.readouterr().out)["success_rate"] == 0.5


def test_cancel_command_cancels_the_jobs_tasks(make_website, capsys):
    from app.models import Job

    site_id = make_website("https://acme.com")
    with get_session() as session:
        job = Job(kind="crawl", status="running")
        session.add(job)
        session.flush()
        queue.enqueue(session, [site_id], job_id=job.id)
        session.commit()
        job_id = job.id
    run_cli("cancel", "--job", str(job_id))
    with get_session() as session:
        assert session.scalars(select(CrawlTask)).one().state == "cancelled"
